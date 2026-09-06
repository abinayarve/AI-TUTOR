"""
Real RAG Integration Test

Tests:
1. Create a sample PDF
2. Create a sample PPTX
3. Load both through RAGEngine
4. Chunk and embed them
5. Store them in Chroma
6. Retrieve relevant information
7. Verify source/page/slide metadata
8. Verify grounded context
"""

from pathlib import Path
import shutil

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from pptx import Presentation
from pptx.util import Inches

from rag_engine import RAGEngine


# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent
TEST_DATA_DIR = BASE_DIR / "data" / "rag_integration_test"
TEST_DB_DIR = TEST_DATA_DIR / "chroma_db"

PDF_PATH = TEST_DATA_DIR / "ohms_law_test.pdf"
PPTX_PATH = TEST_DATA_DIR / "ohms_law_test.pptx"


# ---------------------------------------------------------------------------
# Create test PDF
# ---------------------------------------------------------------------------

def create_test_pdf():
    print("\nCreating sample PDF...")

    TEST_DATA_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    pdf = canvas.Canvas(
        str(PDF_PATH),
        pagesize=A4,
    )

    pdf.setFont(
        "Helvetica",
        16,
    )

    pdf.drawString(
        60,
        780,
        "Ohm's Law",
    )

    pdf.setFont(
        "Helvetica",
        11,
    )

    lines = [
        "Ohm's Law describes the relationship between voltage,",
        "current, and resistance in an electrical circuit.",
        "",
        "The formula is V = I × R.",
        "",
        "V represents voltage measured in volts.",
        "I represents current measured in amperes.",
        "R represents resistance measured in ohms.",
        "",
        "For example, if a circuit has a current of 2 amperes",
        "and a resistance of 5 ohms, the voltage is 10 volts.",
        "",
        "A common mistake is confusing current with voltage.",
        "Current represents the flow of electric charge, while",
        "voltage represents the potential difference.",
    ]

    y = 740

    for line in lines:
        pdf.drawString(
            60,
            y,
            line,
        )

        y -= 20

    pdf.save()

    print(
        f"Created: {PDF_PATH}"
    )


# ---------------------------------------------------------------------------
# Create test PPTX
# ---------------------------------------------------------------------------

def create_test_pptx():
    print("\nCreating sample PPTX...")

    presentation = Presentation()

    # -----------------------------------------------------------------------
    # Slide 1
    # -----------------------------------------------------------------------

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[1]
    )

    slide.shapes.title.text = (
        "Ohm's Law — Circuit Fundamentals"
    )

    slide.placeholders[1].text = (
        "Ohm's Law connects voltage, current, and resistance.\n"
        "Formula: V = I × R\n"
        "Voltage is measured in volts.\n"
        "Current is measured in amperes.\n"
        "Resistance is measured in ohms."
    )

    # -----------------------------------------------------------------------
    # Slide 2
    # -----------------------------------------------------------------------

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[1]
    )

    slide.shapes.title.text = (
        "Worked Example"
    )

    slide.placeholders[1].text = (
        "Given current I = 3 A and resistance R = 4 Ω.\n"
        "Using V = I × R,\n"
        "V = 3 × 4 = 12 V.\n"
        "Therefore, the voltage is 12 volts."
    )

    # -----------------------------------------------------------------------
    # Slide 3
    # -----------------------------------------------------------------------

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[1]
    )

    slide.shapes.title.text = (
        "Common Misconception"
    )

    slide.placeholders[1].text = (
        "Voltage and current are not the same quantity.\n"
        "Voltage is potential difference.\n"
        "Current is the flow of electric charge.\n"
        "Resistance opposes the flow of current."
    )

    presentation.save(
        str(PPTX_PATH)
    )

    print(
        f"Created: {PPTX_PATH}"
    )


# ---------------------------------------------------------------------------
# Test document extraction
# ---------------------------------------------------------------------------

def test_pdf_extraction(engine):
    print("\n" + "=" * 70)
    print("TEST 1: PDF extraction")
    print("=" * 70)

    documents = engine.load_file(
        str(PDF_PATH)
    )

    assert len(documents) > 0

    print(
        f"Extracted {len(documents)} PDF document section(s)"
    )

    first = documents[0]

    print(
        "Source:",
        first.metadata.get("filename"),
    )

    print(
        "Location:",
        first.metadata.get("location"),
    )

    assert (
        first.metadata.get("file_type")
        == "pdf"
    )

    assert (
        first.metadata.get("page")
        == 1
    )

    assert "Ohm's Law" in first.page_content

    print("PASS")


# ---------------------------------------------------------------------------
# Test PPTX extraction
# ---------------------------------------------------------------------------

def test_pptx_extraction(engine):
    print("\n" + "=" * 70)
    print("TEST 2: PPTX extraction")
    print("=" * 70)

    documents = engine.load_file(
        str(PPTX_PATH)
    )

    assert len(documents) == 3

    print(
        f"Extracted {len(documents)} slides"
    )

    for document in documents:
        print(
            f"  {document.metadata.get('location')}: "
            f"{document.metadata.get('filename')}"
        )

    assert (
        documents[0].metadata.get("slide")
        == 1
    )

    assert (
        documents[1].metadata.get("slide")
        == 2
    )

    assert (
        documents[2].metadata.get("slide")
        == 3
    )

    assert (
        "V = I"
        in documents[0].page_content
    )

    assert (
        "Worked Example"
        in documents[1].page_content
    )

    print("PASS")


# ---------------------------------------------------------------------------
# Test PDF ingestion
# ---------------------------------------------------------------------------

def test_pdf_ingestion(engine):
    print("\n" + "=" * 70)
    print("TEST 3: PDF ingestion into Chroma")
    print("=" * 70)

    result = engine.ingest_file(
        str(PDF_PATH),
        clear_existing=True,
    )

    print(
        "Ingestion result:",
        result,
    )

    assert result["success"] is True
    assert result["documents"] > 0
    assert result["chunks"] > 0

    print("PASS")


# ---------------------------------------------------------------------------
# Test PPTX ingestion
# ---------------------------------------------------------------------------

def test_pptx_ingestion(engine):
    print("\n" + "=" * 70)
    print("TEST 4: PPTX ingestion into Chroma")
    print("=" * 70)

    result = engine.ingest_file(
        str(PPTX_PATH),
        clear_existing=False,
    )

    print(
        "Ingestion result:",
        result,
    )

    assert result["success"] is True
    assert result["documents"] == 3
    assert result["chunks"] > 0

    print("PASS")


# ---------------------------------------------------------------------------
# Test retrieval from PDF/PPTX
# ---------------------------------------------------------------------------

def test_retrieval(engine):
    print("\n" + "=" * 70)
    print("TEST 5: Semantic retrieval")
    print("=" * 70)

    query = (
        "What is the formula relating voltage, "
        "current and resistance?"
    )

    results = engine.retrieve(
        query,
        top_k=5,
    )

    assert len(results) > 0

    print(
        f"Retrieved {len(results)} relevant chunks"
    )

    for index, document in enumerate(
        results,
        start=1,
    ):
        print(
            f"\nResult {index}"
        )

        print(
            "Source:",
            document.metadata.get(
                "filename"
            ),
        )

        print(
            "Location:",
            document.metadata.get(
                "location"
            ),
        )

        print(
            "Content:",
            document.page_content[:250]
            .replace("\n", " "),
        )

    combined_text = "\n".join(
        document.page_content
        for document in results
    )

    assert (
        "V = I"
        in combined_text
        or "voltage"
        in combined_text.lower()
    )

    print("\nPASS")


# ---------------------------------------------------------------------------
# Test source references
# ---------------------------------------------------------------------------

def test_source_references(engine):
    print("\n" + "=" * 70)
    print("TEST 6: Source references")
    print("=" * 70)

    results = engine.retrieve(
        "worked example of voltage current resistance",
        top_k=5,
    )

    references = engine.get_source_references(
        results
    )

    assert len(references) > 0

    for reference in references:
        print(
            f"- {reference['filename']} "
            f"({reference['location']})"
        )

    print("PASS")


# ---------------------------------------------------------------------------
# Test grounded context
# ---------------------------------------------------------------------------

def test_grounded_context(engine):
    print("\n" + "=" * 70)
    print("TEST 7: Grounded context")
    print("=" * 70)

    result = engine.retrieve_grounded(
        "Explain the common misconception about voltage and current",
        top_k=5,
    )

    assert result["count"] > 0

    context = result["context"]

    print("\nGenerated grounded context:\n")
    print(context[:2500])

    assert "[SOURCE 1:" in context
    assert "[END SOURCE 1]" in context

    assert (
        "voltage"
        in context.lower()
    )

    print("\nPASS")


# ---------------------------------------------------------------------------
# Test collection information
# ---------------------------------------------------------------------------

def test_collection_info(engine):
    print("\n" + "=" * 70)
    print("TEST 8: Chroma collection information")
    print("=" * 70)

    info = engine.get_collection_info()

    print(
        "Collection:",
        info.get("collection"),
    )

    print(
        "Chunks:",
        info.get("chunks"),
    )

    assert (
        info.get("chunks", 0)
        > 0
    )

    print("PASS")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    print("=" * 70)
    print("AI TEACHER — REAL RAG INTEGRATION TEST")
    print("=" * 70)

    # Remove old integration-test data.
    if TEST_DATA_DIR.exists():
        print(
            "\nRemoving previous test data..."
        )

        shutil.rmtree(
            TEST_DATA_DIR,
            ignore_errors=True,
        )

    TEST_DATA_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    # Create sample learning materials.
    create_test_pdf()
    create_test_pptx()

    # Initialize RAG engine.
    engine = RAGEngine(
        persist_directory=str(
            TEST_DB_DIR
        )
    )

    # Run tests.
    test_pdf_extraction(
        engine
    )

    test_pptx_extraction(
        engine
    )

    test_pdf_ingestion(
        engine
    )

    test_pptx_ingestion(
        engine
    )

    test_retrieval(
        engine
    )

    test_source_references(
        engine
    )

    test_grounded_context(
        engine
    )

    test_collection_info(
        engine
    )

    print("\n" + "=" * 70)
    print("ALL REAL RAG INTEGRATION TESTS PASSED")
    print("=" * 70)

    print(
        "\nYour RAG pipeline successfully supports:"
    )

    print("  ✓ PDF extraction")
    print("  ✓ PPTX extraction")
    print("  ✓ Page metadata")
    print("  ✓ Slide metadata")
    print("  ✓ Semantic chunking")
    print("  ✓ Embedding generation")
    print("  ✓ Chroma vector storage")
    print("  ✓ Semantic retrieval")
    print("  ✓ Source references")
    print("  ✓ Grounded context generation")


if __name__ == "__main__":
    main()